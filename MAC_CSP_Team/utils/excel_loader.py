"""Fast, bounded Excel readers for the Country Sweet inventory workbook.

The source workbook contains many calculation and lookup sheets. The first
screen only needs two small model sheets, so this module reads those sheets in
read-only mode with explicit ranges and leaves every other sheet deferred.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from hashlib import sha256
from io import BytesIO
import logging
import re
from time import perf_counter
from typing import Any, Callable, Iterable
import warnings

import openpyxl
import pandas as pd


LOGGER = logging.getLogger("country_sweet.excel")

CANONICAL_COLUMNS = [
    "product", "variety", "classification", "size", "packaging_type", "units",
    "stored_lbs", "forecast_demand_lbs", "harvest_lbs", "sales_lbs", "packout_pct",
    "purchased_loads", "remaining_loads", "year", "week", "reporting_date",
]

# These are deliberately bounded. They cover the real model layouts without
# trusting a worksheet's potentially inflated max_row/max_column metadata.
CURRENT_WEEK_BOUNDS = (1, 60, 1, 34)       # A1:AH60
FORECAST_BOUNDS = (1, 70, 1, 74)           # A1:BV70
INITIAL_SHEETS = ("Current Week Model", "Forecasted Inventory Levels")

ALIASES: dict[str, list[str]] = {
    "product": ["product", "product name", "item", "item name", "description", "sku", "product description"],
    "variety": ["variety", "type", "sweet potato variety", "commodity"],
    "classification": ["classification", "class", "organic conventional", "organic/conventional", "production type", "organic type"],
    "size": ["size", "pack size", "grade", "pack", "package size", "size grade"],
    "packaging_type": ["packaging type", "packaging", "package type", "package", "pack style", "pack description"],
    "units": ["units", "unit", "uom", "unit of measure", "measure"],
    "stored_lbs": ["stored lbs", "stored_lbs", "inventory lbs", "inventory pounds", "inventory", "on hand lbs", "on hand", "ending inventory", "forecasted inventory", "pounds in storage", "storage lbs"],
    "forecast_demand_lbs": ["forecast demand lbs", "forecast_demand_lbs", "forecast demand", "demand lbs", "sales forecast", "demand", "planned sales lbs"],
    "harvest_lbs": ["harvest lbs", "harvest_lbs", "harvest pounds", "harvest", "forecast harvest", "harvest plan"],
    "sales_lbs": ["sales lbs", "sales_lbs", "sales pounds", "sales", "shipments", "shipped lbs", "planned sales"],
    "packout_pct": ["packout pct", "packout_pct", "packout %", "packout percentage", "packout"],
    "purchased_loads": ["purchased loads", "purchased_loads", "loads purchased", "purchases", "purchased"],
    "remaining_loads": ["remaining loads", "remaining_loads", "loads remaining", "open loads", "remaining"],
    "year": ["year", "crop year", "crop_year", "reporting year"],
    "week": ["week", "week number", "week #", "wk", "reporting week", "week ending"],
    "reporting_date": ["reporting date", "report date", "date", "week ending date", "period", "as of date"],
}


@dataclass
class SheetReport:
    name: str
    status: str
    semantic: str | None = None
    header_row: int | None = None
    row_count: int = 0
    columns: list[str] = field(default_factory=list)
    missing_values: dict[str, int] = field(default_factory=dict)
    formula_cells: int = 0
    formula_errors: int = 0
    missing_formula_results: int = 0
    elapsed_ms: float = 0.0
    frame: pd.DataFrame | None = None


@dataclass
class WorkbookReport:
    available_sheets: list[str]
    preferred_sheet: str
    frames: dict[str, pd.DataFrame]
    sheet_status: dict[str, str]
    warnings: list[str]
    sheet_reports: dict[str, SheetReport] = field(default_factory=dict)
    semantic_sheets: dict[str, str] = field(default_factory=dict)
    timings: dict[str, float] = field(default_factory=dict)
    content_id: str = ""
    deferred_sheets: list[str] = field(default_factory=list)


def workbook_content_id(content: bytes) -> str:
    """Return a stable cache key for the uploaded workbook bytes."""
    return sha256(content).hexdigest()


def _norm(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value).strip().lower()).strip()


def _canonical_empty() -> pd.DataFrame:
    return pd.DataFrame(columns=CANONICAL_COLUMNS)


def _quality_counts(frame: pd.DataFrame) -> dict[str, int]:
    return {str(column): int(frame[column].isna().sum()) for column in frame.columns if frame[column].isna().any()}


def normalize_frame(frame: pd.DataFrame) -> pd.DataFrame:
    """Coerce imported fields while preserving missing values and errors."""
    result = frame.copy()
    for field_name in ["product", "variety", "classification", "size", "packaging_type", "units", "inventory_state"]:
        if field_name in result:
            result[field_name] = result[field_name].where(result[field_name].notna(), pd.NA).astype("string").str.strip()
    if "classification" in result:
        result["classification"] = result["classification"].str.title()
    for field_name in ["stored_lbs", "forecast_demand_lbs", "harvest_lbs", "sales_lbs", "packout_pct", "purchased_loads", "remaining_loads", "year", "week"]:
        if field_name in result:
            result[field_name] = pd.to_numeric(result[field_name], errors="coerce")
    if "reporting_date" in result:
        result["reporting_date"] = pd.to_datetime(result["reporting_date"], errors="coerce")
    return result


def _header_map(columns: Iterable[object]) -> dict[str, object]:
    normalized = {_norm(c): c for c in columns if c is not None and str(c).strip()}
    result: dict[str, object] = {}
    for target, aliases in ALIASES.items():
        for alias in aliases:
            if _norm(alias) in normalized:
                result[target] = normalized[_norm(alias)]
                break
    return result


def _rows(ws: Any, bounds: tuple[int, int, int, int]) -> list[tuple[Any, ...]]:
    min_row, max_row, min_col, max_col = bounds
    return list(ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col, values_only=True))


def _formula_quality_bounded(formula_ws: Any, cached_ws: Any, bounds: tuple[int, int, int, int]) -> tuple[int, int, int]:
    """Inspect formulas only inside a target sheet's bounded model range."""
    formula_cells = formula_errors = missing_results = 0
    min_row, max_row, min_col, max_col = bounds
    formula_rows = formula_ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col, values_only=True)
    cached_rows = cached_ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col, values_only=True)
    for formula_row, cached_row in zip(formula_rows, cached_rows):
        for formula, cached in zip(formula_row, cached_row):
            if not (isinstance(formula, str) and formula.startswith("=")):
                continue
            formula_cells += 1
            if isinstance(cached, str) and cached.startswith("#"):
                formula_errors += 1
            elif cached is None:
                missing_results += 1
    return formula_cells, formula_errors, missing_results


_SIZE_PATTERNS = ("10-12 oz", "10-12oz", "#1 Jewel", "Medium", "Jumbo", "14oz", "Baby", "Small", "#1")


def _split_product_label(label: str, classification: str | None) -> tuple[str, str | None, str | None]:
    clean = str(label).strip()
    prefix = "Conv. " if clean.startswith("Conv. ") else "Org. " if clean.startswith("Org. ") else ""
    inferred_class = classification or ("Conventional" if prefix == "Conv. " else "Organic" if prefix == "Org. " else None)
    base = clean[len(prefix):].strip() if prefix else clean
    for suffix in _SIZE_PATTERNS:
        if base.lower().endswith(suffix.lower()):
            variety = base[: -len(suffix)].strip()
            return clean, variety or None, suffix
    return clean, base or None, None


def _parse_forecasted_inventory(rows: list[tuple[Any, ...]]) -> tuple[pd.DataFrame, dict[str, Any]]:
    """Parse the forecast matrix from F:BV without touching other sheets."""
    if len(rows) < 9:
        return _canonical_empty(), {"actual_date": None, "columns": 0}
    state_row, week_row, label_row, date_row = rows[2], rows[3], rows[4], rows[5]
    timeline_columns = [index for index, value in enumerate(date_row) if value is not None and hasattr(value, "year") and index >= 7]
    actual_dates = [date_row[index] for index in timeline_columns if str(state_row[index]).upper() == "ACTUAL"]
    actual_date = max(actual_dates) if actual_dates else None
    output: list[dict[str, Any]] = []
    classification: str | None = None
    for row_index, row in enumerate(rows[6:51], start=6):
        first = row[5] if len(row) > 5 else None
        description = row[6] if len(row) > 6 else None
        if isinstance(description, str) and description.strip() in {"Conventional", "Organic"}:
            classification = description.strip()
            continue
        if row_index < 8:
            continue
        if first in (None, "Lookups", "Total") and not (isinstance(description, str) and description.startswith(("CSP", "Sweet Potatoes"))):
            continue
        product_label = str(first or description or "").strip()
        if not product_label or product_label in {"Bulk Variety/Size", "Total"}:
            continue
        product, variety, size = _split_product_label(product_label, classification)
        for index in timeline_columns:
            value = row[index] if index < len(row) else None
            if value is None:
                continue
            report_date = date_row[index]
            output.append({
                "product": product, "variety": variety, "classification": classification, "size": size,
                "packaging_type": "Bulk Bin", "units": "lbs", "stored_lbs": value,
                "forecast_demand_lbs": pd.NA, "harvest_lbs": pd.NA, "sales_lbs": pd.NA, "packout_pct": pd.NA,
                "purchased_loads": pd.NA, "remaining_loads": pd.NA,
                "year": report_date.year if hasattr(report_date, "year") else pd.NA,
                "week": week_row[index], "reporting_date": report_date,
                "inventory_state": str(state_row[index]).title() if state_row[index] is not None else pd.NA,
            })
    return normalize_frame(pd.DataFrame(output)), {"actual_date": actual_date, "columns": len(timeline_columns)}


def _parse_current_week(rows: list[tuple[Any, ...]], year_hint: int | None, reporting_date: Any) -> pd.DataFrame:
    """Parse the current-week snapshot from the fixed model table D:AH."""
    if len(rows) < 13:
        return _canonical_empty()
    week_value = rows[5][6] if len(rows[5]) > 6 else pd.NA
    output: list[dict[str, Any]] = []
    for row in rows[12:55]:
        product_source = row[4] if len(row) > 4 else None
        fallback_product = row[3] if len(row) > 3 else None
        product_label = str(product_source or fallback_product or "").strip()
        if not product_label or product_label in {"Total", "Lookups"}:
            continue
        variety_source = row[5] if len(row) > 5 else None
        size = row[6] if len(row) > 6 else None
        classification = "Conventional" if product_label.startswith("Conv. ") else "Organic" if product_label.startswith("Org. ") else None
        _, inferred_variety, inferred_size = _split_product_label(product_label, classification)
        variety = str(variety_source).strip() if variety_source is not None else inferred_variety
        if variety.startswith("Conv. ") or variety.startswith("Org. "):
            variety = variety.split(". ", 1)[1]
        output.append({
            "product": product_label, "variety": variety, "classification": classification,
            "size": str(size).strip() if size is not None else inferred_size,
            "packaging_type": "Bulk Bin", "units": "lbs",
            "stored_lbs": row[28] if len(row) > 28 else pd.NA,
            "forecast_demand_lbs": row[7] if len(row) > 7 else pd.NA,
            "harvest_lbs": pd.NA, "sales_lbs": pd.NA, "packout_pct": pd.NA,
            "purchased_loads": pd.NA, "remaining_loads": pd.NA,
            "year": year_hint if year_hint is not None else pd.NA, "week": week_value,
            "reporting_date": reporting_date, "inventory_state": "Current",
        })
    return normalize_frame(pd.DataFrame(output))


def _report_for_frame(name: str, status: str, semantic: str, frame: pd.DataFrame, elapsed_ms: float, formula_quality: tuple[int, int, int]) -> SheetReport:
    formula_cells, formula_errors, missing_results = formula_quality
    return SheetReport(name=name, status=status, semantic=semantic, row_count=len(frame), columns=[str(column) for column in frame.columns], missing_values=_quality_counts(frame), formula_cells=formula_cells, formula_errors=formula_errors, missing_formula_results=missing_results, elapsed_ms=elapsed_ms, frame=frame)


def load_inventory_workbook(content: bytes, progress_callback: Callable[[str], None] | None = None) -> WorkbookReport:
    """Recognize the workbook and parse only the two initial model sheets."""
    total_started = perf_counter()
    content_id = workbook_content_id(content)
    timings: dict[str, float] = {}

    def stage(message: str) -> None:
        LOGGER.info("excel.stage %s", message)
        if progress_callback:
            progress_callback(message)

    stage("Opening workbook in read-only mode")
    open_started = perf_counter()
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Slicer List extension is not supported and will be removed")
        values_wb = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)
    timings["workbook_open_ms"] = (perf_counter() - open_started) * 1000
    available = list(values_wb.sheetnames)
    timings["sheet_recognition_ms"] = (perf_counter() - open_started) * 1000
    LOGGER.info("excel.sheet_recognition sheets=%d elapsed_ms=%.1f", len(available), timings["sheet_recognition_ms"])

    # Formula inspection is also bounded to the two target ranges.
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Slicer List extension is not supported and will be removed")
        formula_wb = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=False)
    frames: dict[str, pd.DataFrame] = {}
    reports: dict[str, SheetReport] = {}
    statuses: dict[str, str] = {}
    semantic_sheets: dict[str, str] = {}
    quality_warnings: list[str] = []
    deferred = [sheet for sheet in available if sheet not in INITIAL_SHEETS]
    for sheet in deferred:
        reports[sheet] = SheetReport(sheet, "Deferred until this worksheet is explicitly requested.", semantic=None)
        statuses[sheet] = reports[sheet].status

    forecast_meta: dict[str, Any] = {"actual_date": None, "columns": 0}
    parse_started = perf_counter()
    # Forecast first so Current Week Model can inherit its verified year/date.
    for sheet in ("Forecasted Inventory Levels", "Current Week Model"):
        if sheet not in available:
            continue
        stage(f"Parsing {sheet} using a bounded model range")
        sheet_started = perf_counter()
        bounds = FORECAST_BOUNDS if sheet == "Forecasted Inventory Levels" else CURRENT_WEEK_BOUNDS
        values_rows = _rows(values_wb[sheet], bounds)
        formula_quality = _formula_quality_bounded(formula_wb[sheet], values_wb[sheet], bounds)
        if sheet == "Forecasted Inventory Levels":
            frame, forecast_meta = _parse_forecasted_inventory(values_rows)
            semantic = "inventory_timeline"
        else:
            year_hint = int(forecast_meta["actual_date"].year) if forecast_meta.get("actual_date") else None
            frame = _parse_current_week(values_rows, year_hint, forecast_meta.get("actual_date"))
            semantic = "inventory"
        elapsed_ms = (perf_counter() - sheet_started) * 1000
        frames[sheet] = frame
        status = f"Imported bounded range ({len(frame):,} rows) in {elapsed_ms:.1f} ms."
        reports[sheet] = _report_for_frame(sheet, status, semantic, frame, elapsed_ms, formula_quality)
        statuses[sheet] = status
        semantic_sheets[semantic] = sheet
        if formula_quality[1]:
            quality_warnings.append(f"{sheet}: {formula_quality[1]} spreadsheet formula error(s) were found; values remain flagged.")
        if formula_quality[2]:
            quality_warnings.append(f"{sheet}: {formula_quality[2]} formula result(s) have no cached value; they were not treated as zero.")
        LOGGER.info("excel.sheet_parse sheet=%s rows=%d elapsed_ms=%.1f", sheet, len(frame), elapsed_ms)
    timings["target_sheet_parsing_ms"] = (perf_counter() - parse_started) * 1000

    validation_started = perf_counter()
    if "inventory" not in semantic_sheets:
        quality_warnings.append("Current Week Model was not available, so current inventory cards remain unavailable.")
    if "inventory_timeline" not in semantic_sheets:
        quality_warnings.append("Forecasted Inventory Levels was not available, so the inventory timeline remains unavailable.")
    timings["validation_ms"] = (perf_counter() - validation_started) * 1000
    timings["total_ms"] = (perf_counter() - total_started) * 1000
    LOGGER.info("excel.validation elapsed_ms=%.1f", timings["validation_ms"])
    LOGGER.info("excel.total content_id=%s elapsed_ms=%.1f deferred=%d", content_id[:12], timings["total_ms"], len(deferred))
    values_wb.close()
    formula_wb.close()

    return WorkbookReport(
        available_sheets=available,
        preferred_sheet=semantic_sheets.get("inventory", ""),
        frames=frames,
        sheet_status=statuses,
        warnings=quality_warnings,
        sheet_reports=reports,
        semantic_sheets=semantic_sheets,
        timings=timings,
        content_id=content_id,
        deferred_sheets=deferred,
    )


def frame_for(report: WorkbookReport, semantic: str) -> tuple[str | None, pd.DataFrame | None]:
    sheet = report.semantic_sheets.get(semantic)
    if not sheet or sheet not in report.frames:
        return None, None
    return sheet, report.frames[sheet]


def report_to_frame(report: WorkbookReport) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    for item in report.sheet_reports.values():
        rows.append({
            "Worksheet": item.name, "Status": item.status, "Use": item.semantic or "Deferred",
            "Rows": item.row_count, "Elapsed ms": round(item.elapsed_ms, 1),
            "Formula cells": item.formula_cells, "Formula errors": item.formula_errors,
            "Missing cached results": item.missing_formula_results,
            "Missing values": ", ".join(f"{field} ({count})" for field, count in item.missing_values.items()) or "None",
        })
    return pd.DataFrame(rows)


def timings_to_frame(report: WorkbookReport) -> pd.DataFrame:
    return pd.DataFrame([{"Stage": key.replace("_", " ").title(), "Elapsed ms": round(value, 1)} for key, value in report.timings.items()])
